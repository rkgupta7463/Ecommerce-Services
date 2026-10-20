from fastapi import HTTPException,status
from io import BytesIO
from openpyxl import load_workbook
from query.category import category_create
from utility import send_email_background

async def cat_bulk_creation(file,user):
    try:
            # Verify file extension
        if not file.filename.endswith(('.xlsx', '.xls')):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, 
                detail="Invalid file type. Only Excel files are allowed."
            )
        content=await file.read()
        if not content:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded file is empty."
            )
        
        # 3. Load Excel workbook from memory
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True
        )

        worksheet = workbook.active

        # 4. Read headers
        headers = [
            str(cell).strip().lower() if cell is not None else ""
            for cell in next(worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True
            ))
        ]

        required_columns = {
            "name",
            "description",
            "parent_cat_id",
            "is_active"
        }

        missing_columns = required_columns - set(headers)

        if missing_columns:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing columns: {sorted(missing_columns)}"
            )

        categories=[]

        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2
        ):
            record = dict(zip(headers, row))

            # Skip completely empty rows
            if not any(value is not None for value in row):
                continue

            name = record.get("name")

            if not isinstance(name, str) or not name.strip():
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Row {row_number}: category name is required."
                )

            categories.append({
                "name": name.strip(),
                "description": record.get("description"),
                "parent_cat_id": record.get("parent_cat_id"),
                "is_active": record.get("is_active", True)
            })

        workbook.close()

        if not categories:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No category records found in the Excel file."
            )

        # print("categories:- \n",categories)

        
        for cat in categories:
            # category_create(cat)
            cat['user_id']=user
            category_create(data=cat)

        data = {
            "message": "Excel file processed successfully.",
            "total_categories": len(categories),
            "categories": categories
        }

        category_names = "\n".join(
            f"- {category['name']}"
            for category in categories
        )

        email_body = f"""
        Dear Team,

        Your category Excel file has been processed successfully.

        Total categories processed: {len(categories)}

        Category list:
        {category_names}

        Regards,
        E-commerce System
        """

        send_email_background(
            recipient="rishukumr0187@gmail.com",
            subject="Category Bulk Upload - Processing Summary",
            body=email_body
        )

        return data
    
    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unable to process Excel file: {str(e)}"
        )